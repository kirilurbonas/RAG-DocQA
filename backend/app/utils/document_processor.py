"""Document processing utilities for extracting text from various file formats."""

import os
import re
import zipfile
from typing import List, Dict, Optional
from pathlib import Path
from io import BytesIO
import PyPDF2
import docx
import markdown

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None


MAX_PDF_PAGES = int(os.getenv("MAX_PDF_PAGES", "500"))
MAX_EXTRACTED_TEXT_CHARS = int(os.getenv("MAX_EXTRACTED_TEXT_CHARS", str(5_000_000)))


class DocumentProcessor:
    """Processes documents in various formats and extracts text."""

    # .doc is legacy binary; only .docx is safely supported for Word
    SUPPORTED_EXTENSIONS = {'.pdf', '.docx', '.md', '.txt', '.markdown', '.xlsx', '.pptx'}
    
    @staticmethod
    def is_supported(filename: str) -> bool:
        """Check if file extension is supported."""
        ext = Path(filename).suffix.lower()
        return ext in DocumentProcessor.SUPPORTED_EXTENSIONS
    
    @staticmethod
    def extract_text(file_content: bytes, filename: str) -> str:
        """Extract text from file content based on file extension."""
        ext = Path(filename).suffix.lower()

        if ext == '.pdf':
            return DocumentProcessor._extract_from_pdf(file_content)
        elif ext == '.docx':
            return DocumentProcessor._extract_from_docx(file_content)
        elif ext in ['.md', '.markdown']:
            return DocumentProcessor._extract_from_markdown(file_content)
        elif ext == '.xlsx':
            return DocumentProcessor._extract_from_xlsx(file_content)
        elif ext == '.pptx':
            return DocumentProcessor._extract_from_pptx(file_content)
        elif ext == '.txt':
            raw = file_content.decode('utf-8', errors='replace')
            if len(raw) > MAX_EXTRACTED_TEXT_CHARS:
                return raw[:MAX_EXTRACTED_TEXT_CHARS]
            return raw
        else:
            raise ValueError(f"Unsupported file type: {ext}")
    
    @staticmethod
    def _extract_from_pdf(content: bytes) -> str:
        """Extract text from PDF file."""
        pdf_file = BytesIO(content)
        pdf_reader = PyPDF2.PdfReader(pdf_file)
        n_pages = len(pdf_reader.pages)
        if n_pages > MAX_PDF_PAGES:
            raise ValueError(
                f"PDF has too many pages ({n_pages}); maximum allowed is {MAX_PDF_PAGES}"
            )
        text = []
        total_chars = 0
        for page in pdf_reader.pages:
            chunk = page.extract_text() or ""
            if total_chars + len(chunk) > MAX_EXTRACTED_TEXT_CHARS:
                allowed = MAX_EXTRACTED_TEXT_CHARS - total_chars
                if allowed > 0:
                    text.append(chunk[:allowed])
                break
            text.append(chunk)
            total_chars += len(chunk)
        return '\n'.join(text)
    
    @staticmethod
    def _extract_from_docx(content: bytes) -> str:
        """Extract text from DOCX file."""
        doc_file = BytesIO(content)

        # Check for zip bomb/decompression bomb by validating uncompressed size
        # DOCX files are ZIPs - enforce extracted size limits
        try:
            with zipfile.ZipFile(doc_file, 'r') as zip_ref:
                total_uncompressed = sum(info.file_size for info in zip_ref.infolist())
                if total_uncompressed > MAX_EXTRACTED_TEXT_CHARS:
                    raise ValueError(
                        f"DOCX file too large when uncompressed ({total_uncompressed} bytes); "
                        f"maximum allowed is {MAX_EXTRACTED_TEXT_CHARS} bytes"
                    )
        except zipfile.BadZipFile as e:
            raise ValueError(f"Invalid DOCX file format: {e}")

        doc_file.seek(0)  # Reset file pointer after validation
        doc = docx.Document(doc_file)
        paragraphs = [para.text for para in doc.paragraphs]
        out = '\n'.join(paragraphs)
        if len(out) > MAX_EXTRACTED_TEXT_CHARS:
            return out[:MAX_EXTRACTED_TEXT_CHARS]
        return out
    
    @staticmethod
    def _extract_from_markdown(content: bytes) -> str:
        """Extract text from Markdown file."""
        md_text = content.decode('utf-8')
        if len(md_text) > MAX_EXTRACTED_TEXT_CHARS:
            md_text = md_text[:MAX_EXTRACTED_TEXT_CHARS]
        # Convert markdown to plain text (remove formatting)
        html = markdown.markdown(md_text)
        # Simple HTML tag removal (for basic markdown)
        text = re.sub(r'<[^>]+>', '', html)
        return text

    @staticmethod
    def _extract_from_xlsx(content: bytes) -> str:
        """Extract text from Excel file (.xlsx)."""
        if openpyxl is None:
            raise ValueError("openpyxl is required for Excel file support. Install with: pip install openpyxl")

        try:
            # Check for zip bomb/decompression bomb by validating uncompressed size
            # Excel files are ZIPs - enforce extracted size limits
            excel_file = BytesIO(content)

            # Validate zip structure before loading
            try:
                with zipfile.ZipFile(excel_file, 'r') as zip_ref:
                    total_uncompressed = sum(info.file_size for info in zip_ref.infolist())
                    if total_uncompressed > MAX_EXTRACTED_TEXT_CHARS:
                        raise ValueError(
                            f"Excel file too large when uncompressed ({total_uncompressed} bytes); "
                            f"maximum allowed is {MAX_EXTRACTED_TEXT_CHARS} bytes"
                        )
            except zipfile.BadZipFile as e:
                raise ValueError(f"Invalid Excel file format: {e}")

            excel_file.seek(0)  # Reset file pointer after validation
            workbook = openpyxl.load_workbook(excel_file, data_only=True)

            text_parts = []
            total_chars = 0

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text_parts.append(f"Sheet: {sheet_name}\n")
                total_chars += len(f"Sheet: {sheet_name}\n")

                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(cell or "") for cell in row) + "\n"
                    if total_chars + len(row_text) > MAX_EXTRACTED_TEXT_CHARS:
                        text_parts.append(row_text[:MAX_EXTRACTED_TEXT_CHARS - total_chars])
                        return "".join(text_parts)
                    text_parts.append(row_text)
                    total_chars += len(row_text)

            return "".join(text_parts)
        except Exception as e:
            raise ValueError(f"Failed to extract text from Excel file: {e}")

    @staticmethod
    def _extract_from_pptx(content: bytes) -> str:
        """Extract text from PowerPoint file (.pptx)."""
        if Presentation is None:
            raise ValueError("python-pptx is required for PowerPoint file support. Install with: pip install python-pptx")

        try:
            # Check for zip bomb/decompression bomb by validating uncompressed size
            # PowerPoint files are ZIPs - enforce extracted size limits
            pptx_file = BytesIO(content)

            # Validate zip structure before loading
            try:
                with zipfile.ZipFile(pptx_file, 'r') as zip_ref:
                    total_uncompressed = sum(info.file_size for info in zip_ref.infolist())
                    if total_uncompressed > MAX_EXTRACTED_TEXT_CHARS:
                        raise ValueError(
                            f"PowerPoint file too large when uncompressed ({total_uncompressed} bytes); "
                            f"maximum allowed is {MAX_EXTRACTED_TEXT_CHARS} bytes"
                        )
            except zipfile.BadZipFile as e:
                raise ValueError(f"Invalid PowerPoint file format: {e}")

            pptx_file.seek(0)  # Reset file pointer after validation
            presentation = Presentation(pptx_file)

            text_parts = []
            total_chars = 0

            for slide_num, slide in enumerate(presentation.slides, 1):
                text_parts.append(f"\n--- Slide {slide_num} ---\n")
                total_chars += len(f"\n--- Slide {slide_num} ---\n")

                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        shape_text = shape.text + "\n"
                        if total_chars + len(shape_text) > MAX_EXTRACTED_TEXT_CHARS:
                            text_parts.append(shape_text[:MAX_EXTRACTED_TEXT_CHARS - total_chars])
                            return "".join(text_parts)
                        text_parts.append(shape_text)
                        total_chars += len(shape_text)

            return "".join(text_parts)
        except Exception as e:
            raise ValueError(f"Failed to extract text from PowerPoint file: {e}")

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> List[Dict[str, any]]:
        """
        Split text into chunks with overlap.
        
        Args:
            text: Text to chunk
            chunk_size: Maximum size of each chunk in characters
            chunk_overlap: Number of characters to overlap between chunks
            
        Returns:
            List of dictionaries with 'text' and 'metadata' keys
        """
        if len(text) <= chunk_size:
            return [{'text': text, 'metadata': {'chunk_index': 0, 'start': 0, 'end': len(text)}}]
        
        chunks = []
        start = 0
        chunk_index = 0
        
        while start < len(text):
            end = start + chunk_size
            chunk_text = text[start:end]
            
            # Try to break at sentence boundary
            if end < len(text):
                # Look for sentence endings
                last_period = chunk_text.rfind('.')
                last_newline = chunk_text.rfind('\n')
                break_point = max(last_period, last_newline)
                
                if break_point > chunk_size * 0.5:  # Only break if we're past halfway
                    chunk_text = chunk_text[:break_point + 1]
                    end = start + break_point + 1
            
            chunks.append({
                'text': chunk_text.strip(),
                'metadata': {
                    'chunk_index': chunk_index,
                    'start': start,
                    'end': end
                }
            })
            
            start = end - chunk_overlap
            chunk_index += 1
        
        return chunks
